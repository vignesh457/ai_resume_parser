import io
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.shared import qn

def build_docx_binary(user_data):
    """Compiles structured JSON data into a clean, professional, high-density .docx layout"""
    doc = docx.Document()
    
    # Page Margins: 0.5 inch all around for engineering resumes
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11.0)
        section.top_margin = section.bottom_margin = Inches(0.5)
        section.left_margin = section.right_margin = Inches(0.5)

    # Color Palette definitions
    COLOR_PRIMARY = RGBColor(11, 83, 148)  # Deep Professional Blue
    COLOR_TEXT = RGBColor(0, 0, 0)         # Stark Black
    COLOR_MUTED = RGBColor(60, 60, 60)     # Dark Grey

    def set_run_font(run, name='Calibri', size=10, bold=False, italic=False, color=COLOR_TEXT):
        run.font.name = name
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color

    # Enforce Calibri across common styles so hyperlink runs inherit the font
    try:
        normal_style = doc.styles['Normal']
        normal_style.font.name = 'Calibri'
        normal_style.font.size = Pt(10)
    except Exception:
        pass

    try:
        hyperlink_style = doc.styles['Hyperlink']
        hyperlink_style.font.name = 'Calibri'
        hyperlink_style.font.size = Pt(9.5)
        hyperlink_style.font.color.rgb = COLOR_PRIMARY
    except Exception:
        pass

    # Ensure list and heading styles also use Calibri where available
    for _s in ('List Bullet', 'List Number', 'Heading 1', 'Heading 2'):
        try:
            s = doc.styles[_s]
            s.font.name = 'Calibri'
            s.font.size = Pt(10)
        except Exception:
            pass

    def add_section_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        
        run = p.add_run(text.upper())
        set_run_font(run, name='Calibri', size=10.5, bold=True, color=COLOR_PRIMARY)
        
        # Add clean crisp bottom border line under the section heading
        pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                         r'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="0B5394"/>'
                         r'</w:pBdr>')
        p._p.get_or_add_pPr().append(pBdr)

    def add_hyperlink(paragraph, url, text):
        """Add a clickable hyperlink to a paragraph."""
        part = paragraph.part
        r_id = part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        hyperlink = OxmlElement('w:hyperlink')
        hyperlink.set(qn('r:id'), r_id)

        new_run = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rStyle = OxmlElement('w:rStyle')
        rStyle.set(qn('w:val'), 'Hyperlink')
        rPr.append(rStyle)
        # Specify Calibri font for the hyperlink run so it matches document text
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:ascii'), 'Calibri')
        rFonts.set(qn('w:hAnsi'), 'Calibri')
        rPr.append(rFonts)

        # Set size (10pt -> w:sz value is half-points) and color to primary link color
        rSz = OxmlElement('w:sz')
        rSz.set(qn('w:val'), '20')
        rPr.append(rSz)

        rColor = OxmlElement('w:color')
        rColor.set(qn('w:val'), '0B5394')
        rPr.append(rColor)
        new_run.append(rPr)

        new_text = OxmlElement('w:t')
        new_text.text = text
        new_run.append(new_text)
        hyperlink.append(new_run)

        paragraph._p.append(hyperlink)
        return hyperlink

    def normalize_url(url):
        url_text = str(url or '').strip()
        if not url_text:
            return ''
        if url_text.startswith(('http://', 'https://')):
            return url_text
        return f"https://{url_text}"

    # ==================== HEADER SECTION ====================
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_after = Pt(2)
    run_name = p_name.add_run(user_data.get('name', 'VIGNESHWAR REDDY DONAPATI').upper())
    set_run_font(run_name, name='Calibri', size=18, bold=True, color=COLOR_PRIMARY)

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_after = Pt(2)
    contact_info = f"{user_data.get('phone', '+91 9908731210')} | {user_data.get('email', 'donapativigneshwar@gmail.com')}"
    run_contact = p_contact.add_run(contact_info)
    set_run_font(run_contact, name='Calibri', size=9.5, color=COLOR_TEXT)

    p_links = doc.add_paragraph()
    p_links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_links.paragraph_format.space_after = Pt(6)

    profile_links = [
        (user_data.get('portfolio', 'imvignesh.in'), normalize_url(user_data.get('portfolio', 'imvignesh.in'))),
        (user_data.get('linkedin', 'linkedin.com/in/vignesh457'), normalize_url(user_data.get('linkedin', 'linkedin.com/in/vignesh457'))),
        (user_data.get('github', 'github.com/vignesh457'), normalize_url(user_data.get('github', 'github.com/vignesh457'))),
    ]

    for index, (display_text, url) in enumerate(profile_links):
        if index:
            sep_run = p_links.add_run(" | ")
            set_run_font(sep_run, name='Calibri', size=9.5, color=COLOR_PRIMARY)
        if url:
            add_hyperlink(p_links, url, display_text)
        else:
            run_links = p_links.add_run(display_text)
            set_run_font(run_links, name='Calibri', size=9.5, color=COLOR_PRIMARY)

    # ==================== EXPERIENCE SECTION ====================
    if 'experience' in user_data and user_data['experience']:
        add_section_heading("Experience")
        for job in user_data['experience']:
            p_job = doc.add_paragraph()
            p_job.paragraph_format.space_before = Pt(4)
            p_job.paragraph_format.space_after = Pt(2)
            p_job.paragraph_format.keep_with_next = True
            
            r_title = p_job.add_run(f"{job.get('title', 'Software Engineer')} | {job.get('company', 'Company')} ")
            set_run_font(r_title, name='Calibri', size=10, bold=True, color=COLOR_TEXT)
            
            # Align Dates completely flush to the right margin
            p_job.paragraph_format.tab_stops.add_tab_stop(Inches(7.5), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)
            p_job.add_run("\t")
            r_dates = p_job.add_run(job.get('dates', ''))
            set_run_font(r_dates, name='Calibri', size=10, bold=True, color=COLOR_TEXT)

            for bullet in job.get('bullets', []):
                bp = doc.add_paragraph(style='List Bullet')
                bp.paragraph_format.space_after = Pt(2.5)
                bp.paragraph_format.line_spacing = 1.1
                
                parts = bullet.split("**")
                for idx, part in enumerate(parts):
                    br = bp.add_run(part)
                    set_run_font(br, name='Calibri', size=10, bold=(idx % 2 == 1), color=COLOR_TEXT)

    # ==================== SKILLS SECTION ====================
    if 'skills' in user_data and user_data['skills']:
        add_section_heading("Skills")
        skills_data = user_data['skills']
        if isinstance(skills_data, list):
            p_sk = doc.add_paragraph()
            p_sk.paragraph_format.space_after = Pt(3)
            r_sk = p_sk.add_run(" • ".join(skills_data))
            set_run_font(r_sk, name='Calibri', size=10, color=COLOR_TEXT)
        elif isinstance(skills_data, dict):
            for category, items in skills_data.items():
                p_sk = doc.add_paragraph()
                p_sk.paragraph_format.space_after = Pt(2.5)
                
                cat_run = p_sk.add_run(f"{category}: ")
                set_run_font(cat_run, name='Calibri', size=10, bold=True, color=COLOR_TEXT)
                
                items_str = ", ".join(items) if isinstance(items, list) else str(items)
                item_run = p_sk.add_run(items_str)
                set_run_font(item_run, name='Calibri', size=10, color=COLOR_TEXT)

    # ==================== PROJECTS SECTION ====================
    if 'projects' in user_data and user_data['projects']:
        add_section_heading("Projects")
        for proj in user_data['projects']:
            p_proj = doc.add_paragraph()
            p_proj.paragraph_format.space_before = Pt(4)
            p_proj.paragraph_format.space_after = Pt(1)
            p_proj.paragraph_format.keep_with_next = True
            
            r_pname = p_proj.add_run(f"{proj.get('title', 'Project')} ")
            set_run_font(r_pname, name='Calibri', size=10, bold=True, color=COLOR_TEXT)
            
            if proj.get('link'):
                p_proj.add_run("| ")
                add_hyperlink(p_proj, normalize_url(proj.get('link')), "Link")
                p_proj.add_run(" ")
                
            for bullet in proj.get('bullets', []):
                bp = doc.add_paragraph(style='List Bullet')
                bp.paragraph_format.space_after = Pt(2.5)
                bp.paragraph_format.line_spacing = 1.1
                
                parts = bullet.split("**")
                for idx, part in enumerate(parts):
                    br = bp.add_run(part)
                    set_run_font(br, name='Calibri', size=10, bold=(idx % 2 == 1), color=COLOR_TEXT)

    # ==================== CERTIFICATIONS SECTION ====================
    if 'certifications' in user_data and user_data['certifications']:
        add_section_heading("Certifications")
        for cert in user_data['certifications']:
            p_cert = doc.add_paragraph()
            p_cert.paragraph_format.space_after = Pt(2)
            p_cert.paragraph_format.keep_with_next = True

            title = cert.get('title', '')
            link = cert.get('link', '')
            p_cert.add_run(title)
            if link:
                p_cert.add_run(" | ")
                add_hyperlink(p_cert, normalize_url(link), "Certificate")

    # ==================== CODING PROFILES SECTION ====================
    if 'coding_profiles' in user_data and user_data['coding_profiles']:
        add_section_heading("Coding Profiles")
        profiles = user_data['coding_profiles']
        if isinstance(profiles, list):
            profile_text = " | ".join(str(item) for item in profiles)
        elif isinstance(profiles, dict):
            profile_items = []
            for name, value in profiles.items():
                profile_items.append(f"{name} ({value})")
            profile_text = " | ".join(profile_items)
        else:
            profile_text = str(profiles)

        p_profile = doc.add_paragraph()
        p_profile.paragraph_format.space_after = Pt(3)
        run_profile = p_profile.add_run(profile_text)
        set_run_font(run_profile, name='Calibri', size=10, color=COLOR_TEXT)

    # ==================== EDUCATION SECTION ====================
    if 'education' in user_data and user_data['education']:
        add_section_heading("Education")
        for edu in user_data['education']:
            p_edu = doc.add_paragraph()
            p_edu.paragraph_format.space_after = Pt(2)
            
            r_inst = p_edu.add_run(f"{edu.get('institution', 'University')}\t")
            set_run_font(r_inst, name='Calibri', size=10, bold=True, color=COLOR_TEXT)
            
            p_edu.paragraph_format.tab_stops.add_tab_stop(Inches(7.5), docx.enum.text.WD_TAB_ALIGNMENT.RIGHT)
            r_edates = p_edu.add_run(edu.get('dates', ''))
            set_run_font(r_edates, name='Calibri', size=10, bold=True, color=COLOR_TEXT)
            
            p_det = doc.add_paragraph()
            p_det.paragraph_format.space_after = Pt(3)
            r_det = p_det.add_run(f"{edu.get('degree', '')} in {edu.get('field', '')} | GPA: {edu.get('gpa', '')}")
            set_run_font(r_det, name='Calibri', size=10, italic=False, color=COLOR_MUTED)

    target_stream = io.BytesIO()
    doc.save(target_stream)
    return target_stream.getvalue().hex()